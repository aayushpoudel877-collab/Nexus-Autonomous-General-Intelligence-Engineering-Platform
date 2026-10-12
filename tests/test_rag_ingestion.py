import pytest

from services.rag.ingestion import (
    DocumentExtractionError,
    extract_document,
)


def test_extracts_utf8_text_and_removes_bom():
    result = extract_document(b"\xef\xbb\xbfHello\r\nNexus", filename="notes.txt", content_type="text/plain")
    assert result.text == "Hello\nNexus"
    assert result.parser == "utf-8-text-v1"
    assert result.byte_count == len(b"\xef\xbb\xbfHello\r\nNexus")
    assert result.extracted_chars == len(result.text)


def test_html_extraction_ignores_script_style_and_head_content():
    result = extract_document(
        b"<html><head><title>Hidden title</title></head><body><h1>Course</h1>"
        b"<script>stealSecrets()</script><style>.x{}</style><p>Visible &amp; safe</p></body></html>",
        filename="course.html",
        content_type="text/html",
    )
    assert "Course" in result.text
    assert "Visible & safe" in result.text
    assert "stealSecrets" not in result.text
    assert "Hidden title" not in result.text
    assert ".x" not in result.text


@pytest.mark.parametrize(
    ("filename", "content_type", "data"),
    [
        ("notes.exe", "text/plain", b"hello"),
        ("notes.txt", "application/pdf", b"hello"),
        ("notes.txt", "text/plain", b"\xff"),
        ("empty.html", "text/html", b"<html><script>hidden()</script></html>"),
        ("missing-extension", "text/plain", b"hello"),
    ],
)
def test_rejects_unsupported_mismatched_or_empty_documents(filename, content_type, data):
    with pytest.raises(DocumentExtractionError):
        extract_document(data, filename=filename, content_type=content_type)


def test_rejects_upload_above_configured_byte_limit():
    with pytest.raises(DocumentExtractionError, match="upload exceeds"):
        extract_document(b"12345", filename="notes.txt", max_bytes=4)


def test_rejects_non_pdf_bytes_even_when_named_pdf():
    with pytest.raises(DocumentExtractionError, match="does not match PDF"):
        extract_document(b"not a pdf", filename="notes.pdf", content_type="application/pdf")


def test_sanitizes_path_components_from_client_filename():
    result = extract_document(b"body", filename="../../lecture.txt")
    assert result.filename == "lecture.txt"


def test_rejects_invalid_limit_and_non_bytes():
    with pytest.raises(ValueError):
        extract_document(b"hello", filename="a.txt", max_bytes=0)
    with pytest.raises(TypeError):
        extract_document("hello", filename="a.txt")


def test_extracts_docx_paragraphs_and_tables():
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_paragraph("NEXUS document ingestion")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "retrieval"
    table.cell(0, 1).text = "evaluation"
    buffer = BytesIO()
    document.save(buffer)

    result = extract_document(
        buffer.getvalue(),
        filename="course.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "NEXUS document ingestion" in result.text
    assert "retrieval | evaluation" in result.text
    assert result.parser == "python-docx-text-v1"
