"""Bounded, format-aware document extraction for the RAG ingestion pipeline.

Extracted document content is untrusted data. This module does not execute macros,
scripts, embedded objects, or remote resources.
"""
from __future__ import annotations

from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

SUPPORTED_FORMATS = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 2_000_000
_GENERIC_MIME_TYPES = {"", "application/octet-stream"}


class DocumentExtractionError(ValueError):
    """Raised when a document is unsupported, malformed, or exceeds limits."""


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    filename: str
    media_type: str
    parser: str
    byte_count: int
    extracted_chars: int


class _VisibleHTMLText(HTMLParser):
    """Collect visible HTML text while ignoring active or non-content elements."""

    _IGNORED = {"script", "style", "noscript", "template", "svg", "head"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in self._IGNORED:
            self._ignored_depth += 1
        elif not self._ignored_depth and tag.lower() in {"br", "p", "div", "li", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in self._IGNORED and self._ignored_depth:
            self._ignored_depth -= 1
        elif not self._ignored_depth and tag.lower() in {"p", "div", "li", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth:
            self.parts.append(data)


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError("text documents must use valid UTF-8 encoding") from exc


def _extract_html(data: bytes) -> str:
    parser = _VisibleHTMLText()
    try:
        parser.feed(_decode_text(data))
        parser.close()
    except (UnicodeError, ValueError) as exc:
        raise DocumentExtractionError("HTML document could not be parsed") from exc
    return "".join(parser.parts)


def _extract_pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise DocumentExtractionError("file content does not match PDF format")
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentExtractionError("PDF support is unavailable; install the PDF extra") from exc
    try:
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise DocumentExtractionError("encrypted PDF documents are not supported")
        parts: list[str] = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
            if sum(map(len, parts)) > MAX_EXTRACTED_CHARS:
                raise DocumentExtractionError("extracted text exceeds the supported character limit")
        return "\n".join(parts)
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError("PDF document is malformed or could not be extracted") from exc


def _extract_docx(data: bytes) -> str:
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if "word/document.xml" not in names:
                raise DocumentExtractionError("file content does not match DOCX format")
            info = archive.getinfo("word/document.xml")
            if info.file_size > 20 * 1024 * 1024:
                raise DocumentExtractionError("DOCX document XML exceeds the supported size")
            if sum(entry.file_size for entry in entries) > 50 * 1024 * 1024:
                raise DocumentExtractionError("DOCX expanded archive exceeds the supported size")
    except BadZipFile as exc:
        raise DocumentExtractionError("file content does not match DOCX format") from exc
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentExtractionError("DOCX support is unavailable; install the document extra") from exc
    try:
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise DocumentExtractionError("DOCX document is malformed or could not be extracted") from exc


def extract_document(
    data: bytes,
    *,
    filename: str,
    content_type: str | None = None,
    max_bytes: int = MAX_UPLOAD_BYTES,
) -> ExtractionResult:
    """Validate and extract one bounded document without executing embedded content."""
    if not isinstance(data, bytes):
        raise TypeError("document data must be bytes")
    if not isinstance(filename, str) or not filename.strip():
        raise DocumentExtractionError("filename is required")
    if max_bytes < 1:
        raise ValueError("max_bytes must be positive")
    if len(data) > max_bytes:
        raise DocumentExtractionError(f"upload exceeds the {max_bytes}-byte limit")

    safe_name = PurePosixPath(filename.replace("\\", "/")).name
    suffix = PurePosixPath(safe_name).suffix.lower()
    expected_type = SUPPORTED_FORMATS.get(suffix)
    if expected_type is None:
        raise DocumentExtractionError(f"unsupported file extension: {suffix or '(none)'}")

    declared_type = (content_type or "").split(";", 1)[0].strip().lower()
    if declared_type not in _GENERIC_MIME_TYPES and declared_type != expected_type:
        # Browsers commonly label Markdown as text/plain.
        if not (suffix == ".md" and declared_type == "text/plain"):
            raise DocumentExtractionError("declared content type does not match the file extension")

    if suffix in {".txt", ".md"}:
        text = _decode_text(data)
        parser_name = "utf-8-text-v1"
    elif suffix in {".html", ".htm"}:
        text = _extract_html(data)
        parser_name = "html-visible-text-v1"
    elif suffix == ".pdf":
        text = _extract_pdf(data)
        parser_name = "pypdf-text-v1"
    else:
        text = _extract_docx(data)
        parser_name = "python-docx-text-v1"

    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        raise DocumentExtractionError("document contains no extractable text")
    if len(text) > MAX_EXTRACTED_CHARS:
        raise DocumentExtractionError("extracted text exceeds the supported character limit")
    return ExtractionResult(
        text=text,
        filename=safe_name,
        media_type=expected_type,
        parser=parser_name,
        byte_count=len(data),
        extracted_chars=len(text),
    )
